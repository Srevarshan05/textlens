"""
textlens.documents.office
─────────────────────────
Native extraction for Office documents (``pip install "textlens-ocr[documents]"``).

* **DOCX** — paragraphs with heading levels from paragraph styles, list
  items, and real tables (cells, header row).  Word documents have no fixed
  page geometry, so content is returned as page 1 (explicit page breaks
  start new pages).
* **PPTX** — one page per slide with shape positions in points; tables are
  kept as tables; pictures on a slide are OCR'd through the normal routed
  pipeline unless ``ocr='off'``.
"""

from __future__ import annotations

import io
from typing import Any, Iterator, List

from textlens.core.result import (
    HEADING,
    LIST_ITEM,
    SOURCE_NATIVE,
    TABLE,
    TEXT,
    TITLE,
    BBox,
    Block,
    Line,
    Page,
    PageProvenance,
    Table,
    TableCell,
)
from textlens.documents.layout import finalize_order
from textlens.errors import BackendUnavailableError

_EMU_PER_PT = 12700.0


def _need(module: str, dist: str) -> Any:
    try:
        return __import__(module)
    except ImportError as exc:
        raise BackendUnavailableError(
            f"Reading this file needs {dist}.", hint='pip install "textlens-ocr[documents]"'
        ) from exc


def office_pages(source: Any, pipeline: Any) -> Iterator[Page]:
    if source.kind == "docx":
        yield from _docx_pages(source)
    elif source.kind == "pptx":
        yield from _pptx_pages(source, pipeline)
    else:  # pragma: no cover - sniffing prevents this
        from textlens.errors import UnsupportedInputError

        raise UnsupportedInputError(f"Unsupported document type {source.kind}")


def _block(kind: str, text: str, page: int, level: Any = None, bbox: Any = None, **attrs: Any) -> Block:
    return Block(
        type=kind,
        text=text,
        bbox=bbox,
        confidence=1.0,
        lines=[Line(t, None, 1.0) for t in text.split("\n") if t.strip()],
        page=page,
        source=SOURCE_NATIVE,
        backend="python-docx" if bbox is None else "python-pptx",
        level=level,
        attributes=attrs,
    )


def _docx_pages(source: Any) -> Iterator[Page]:
    docx = _need("docx", "python-docx")
    document = docx.Document(io.BytesIO(source.read_bytes()))
    body = document.element.body
    page_no = 1
    blocks: List[Block] = []
    tables: List[Table] = []

    def flush() -> Page:
        finalize_order(blocks, page_no)
        return Page(
            number=page_no,
            unit="flow",
            blocks=list(blocks),
            tables=list(tables),
            provenance=PageProvenance(page=page_no, source=SOURCE_NATIVE, backend="python-docx", confidence=1.0),
            classification="native",
        )

    para_map = {p._p: p for p in document.paragraphs}
    table_map = {t._tbl: t for t in document.tables}
    for child in body.iterchildren():
        if child in para_map:
            p = para_map[child]
            if 'w:br w:type="page"' in child.xml or "lastRenderedPageBreak" in child.xml and not p.text.strip():
                if blocks or tables:
                    yield flush()
                    page_no += 1
                    blocks, tables = [], []
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            if style == "title":
                blocks.append(_block(TITLE, text, page_no, 1))
            elif style.startswith("heading"):
                digits = "".join(ch for ch in style if ch.isdigit())
                blocks.append(_block(HEADING, text, page_no, int(digits) if digits else 2))
            elif "list" in style or child.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None:
                blocks.append(_block(LIST_ITEM, text, page_no))
            else:
                blocks.append(_block(TEXT, text, page_no))
        elif child in table_map:
            t = table_map[child]
            cells = []
            for r, row in enumerate(t.rows):
                for c, cell in enumerate(row.cells):
                    cells.append(TableCell(row=r, col=c, text=cell.text.strip(), is_header=r == 0, confidence=1.0))
            table = Table(cells=cells, page=page_no, source=SOURCE_NATIVE, confidence=1.0, id=f"p{page_no}-t{len(tables)}")
            tables.append(table)
            blocks.append(_block(TABLE, table.text, page_no, table_index=len(tables) - 1))
    yield flush()


def _pptx_pages(source: Any, pipeline: Any) -> Iterator[Page]:
    pptx = _need("pptx", "python-pptx")
    from pptx.util import Emu  # noqa: F401  (import check)

    prs = pptx.Presentation(io.BytesIO(source.read_bytes()))
    width = (prs.slide_width or 0) / _EMU_PER_PT
    height = (prs.slide_height or 0) / _EMU_PER_PT
    for number, slide in enumerate(prs.slides, start=1):
        if pipeline.options.pages and number not in pipeline.options.pages:
            continue
        blocks: List[Block] = []
        tables: List[Table] = []
        fused = False
        shapes = sorted(slide.shapes, key=lambda s: ((s.top or 0), (s.left or 0)))
        for shape in shapes:
            bbox = None
            if shape.left is not None and shape.top is not None and shape.width is not None:
                x0, y0 = shape.left / _EMU_PER_PT, shape.top / _EMU_PER_PT
                bbox = BBox(x0, y0, x0 + shape.width / _EMU_PER_PT, y0 + shape.height / _EMU_PER_PT).rounded(2)
            if shape.has_text_frame and shape.text_frame.text.strip():
                is_title = shape == getattr(slide.shapes, "title", None)
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                blocks.append(_block(TITLE if is_title else TEXT, text, number, 1 if is_title else None, bbox))
            elif getattr(shape, "has_table", False) and shape.has_table:
                cells = [
                    TableCell(row=r, col=c, text=cell.text.strip(), is_header=r == 0, confidence=1.0)
                    for r, row in enumerate(shape.table.rows)
                    for c, cell in enumerate(row.cells)
                ]
                table = Table(cells=cells, page=number, bbox=bbox, source=SOURCE_NATIVE, confidence=1.0)
                tables.append(table)
                blocks.append(_block(TABLE, table.text, number, None, bbox, table_index=len(tables) - 1))
            elif shape.shape_type == 13 and pipeline.options.ocr != "off":  # MSO_SHAPE_TYPE.PICTURE
                from PIL import Image

                from textlens.core.router import PageSignals

                try:
                    img = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
                except Exception:
                    continue
                if min(img.size) < 32:
                    continue
                sub = pipeline.ocr_image(img, number, PageSignals(kind="figure"))
                if sub.text.strip() and bbox is not None:
                    blk = _block(TEXT, sub.text, number, None, bbox, region="picture")
                    blk.source, blk.model, blk.confidence = "ocr", sub.provenance.model, sub.confidence
                    blocks.append(blk)
                    fused = True
        finalize_order(blocks, number)
        yield Page(
            number=number,
            width=round(width, 2),
            height=round(height, 2),
            unit="pt",
            blocks=blocks,
            tables=tables,
            provenance=PageProvenance(page=number, source="fused" if fused else SOURCE_NATIVE, backend="python-pptx", confidence=1.0),
            classification="native",
        )
